import base64
from copy import copy
from datetime import datetime
import gc
import io
import json
import os
import re
import tempfile
import traceback
import zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from fastapi import FastAPI, File, Form, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pypdf
from starlette.background import BackgroundTask

app = FastAPI(title="QT1 Doc Control API - High Speed Master Engine")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition", "X-Processed-Items"],
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def clean_safe_filename(name: str) -> str:
    cleaned = re.sub(r'[\\/*?:"<>|]', '-', name).strip()
    return cleaned[:130] if len(cleaned) > 130 else cleaned

def normalize_doc_code(raw: str) -> str:
    if not raw: return ""
    m = re.search(r'(QT1-[0-9L]-[A-Za-z0-9_-]+)', raw)
    if not m: return ""
    code = m.group(1).strip()
    code = re.sub(r'_(?:[A-Za-z0-9]{1,3}|Rev\.?[A-Za-z0-9]+)$', '', code, flags=re.IGNORECASE)
    code = re.sub(r'[\s_]+rev\.?.*$', '', code, flags=re.IGNORECASE)
    return code

def clean_document_title(raw_title: str) -> str:
    if not raw_title: return ""
    t = str(raw_title).strip()
    t = re.sub(r'^(?:iv|v|vi|vii|viii|ix|x|\d{1,2})\s+', '', t, flags=re.IGNORECASE).strip()
    t = re.sub(r'\s*\([A-Za-z0-9\s,.~-]*(?:APD|IFA|IFI|AB|RJ|Rev\.?|R\d)[A-Za-z0-9\s,.~-]*\)$', '', t, flags=re.IGNORECASE).strip()
    t = re.sub(r'VoltageTransformer', 'Voltage Transformer', t, flags=re.IGNORECASE)
    if re.match(r'^(?:iv|v|vi|vii|viii|ix|x|\d{1,2})$', t, re.IGNORECASE):
        return ""
    return t

def extract_title_from_filename(filename: str) -> str:
    s = os.path.splitext(filename)[0].strip()
    s = re.sub(r'^(?:\d+|00xx|xxxx|Xxxx)\s*CaRs?\s*R\d+\s*(?:\d{1,2}\s*[A-Za-z]{3}\s*\d{4}\s*)?', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'^(?:QT1-[A-Za-z0-9_-]+)\s*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'[-_](?:reply|comments?|rev\.?[A-Za-z0-9]+)$', '', s, flags=re.IGNORECASE).strip()
    return clean_document_title(s)

def extract_cover_date_bulletproof(text: str) -> str:
    month_map = {
        'jan': 1, 'january': 1, 'feb': 2, 'february': 2, 'mar': 3, 'march': 3,
        'apr': 4, 'april': 4, 'may': 5, 'jun': 6, 'june': 6, 'jul': 7, 'july': 7,
        'aug': 8, 'august': 8, 'sep': 9, 'september': 9, 'sept': 9,
        'oct': 10, 'october': 10, 'nov': 11, 'november': 11, 'dec': 12, 'december': 12
    }
    month_pat = r'(?:January|February|March|April|May|June|July|August|September|October|November|December|Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)'
    
    m_date = re.search(r'\bDate\b', text, re.IGNORECASE)
    if m_date:
        snippet = text[m_date.end(): m_date.end() + 120]
        snippet_clean = re.sub(r'Day\s*Mon\s*Year', '', snippet, flags=re.IGNORECASE)
        snippet_clean = re.sub(r'Sheet\s*:\s*\d+\s*of\s*\d+', '', snippet_clean, flags=re.IGNORECASE)
        
        m_8digits = re.search(r'[:\s-]*\s*(\d)\s*(\d)\s*(\d)\s*(\d)\s*(2)\s*(0)\s*(2)\s*(\d)', snippet_clean)
        if m_8digits:
            g = m_8digits.groups()
            d_val = int(g[0] + g[1])
            m_val = int(g[2] + g[3])
            y_val = g[4] + g[5] + g[6] + g[7]
            if 1 <= m_val <= 12 and 1 <= d_val <= 31:
                return f"{d_val:02d}/{m_val:02d}/{y_val}"

        m_slash = re.search(r'[:\s-]*\s*(\d{1,2})[\/.-](\d{1,2})[\/.-](202\d)', snippet_clean)
        if m_slash:
            d_val = int(m_slash.group(1))
            m_val = int(m_slash.group(2))
            y_val = m_slash.group(3)
            if 1 <= m_val <= 12 and 1 <= d_val <= 31:
                return f"{d_val:02d}/{m_val:02d}/{y_val}"

        collapsed = re.sub(r'\s+', ' ', snippet_clean).strip()
        m_en = re.search(r'\b(' + month_pat + r')\s*(\d{1,2})\s*,?\s*(202\s*\d)', collapsed, re.IGNORECASE)
        if m_en:
            mon_str = m_en.group(1).lower()
            d_val = int(m_en.group(2))
            y_val = re.sub(r'\s+', '', m_en.group(3))
            if mon_str in month_map and 1 <= d_val <= 31:
                return f"{d_val:02d}/{month_map[mon_str]:02d}/{y_val}"

    return ""

def extract_cover_metadata(pdf_bytes: bytes, filename: str = "") -> dict:
    docs = {}
    full_text_pages = []
    
    try:
        reader = pypdf.PdfReader(io.BytesIO(pdf_bytes), strict=False)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                full_text_pages.append(t)
    except Exception as e:
        print(f"pypdf reader error: {e}")
        
    full_text = "\n".join(full_text_pages)

    tr_no = ""
    m_fn = re.search(r'(QT1-[A-Za-z0-9]+-[A-Za-z0-9]+-[A-Za-z0-9]+-[A-Za-z0-9]+)', filename)
    if m_fn:
        tr_no = m_fn.group(1).strip()
    if not tr_no:
        m_ref = re.search(r'(?:TR\.?\s*No\.?|Ref\.?)\s*[:\s]*([A-Za-z0-9_-]*QT1[A-Za-z0-9_-]+)', full_text, re.IGNORECASE)
        if m_ref and "EVNPMB2" not in m_ref.group(1):
            tr_no = m_ref.group(1).strip()
    if not tr_no:
        tr_no = "QT1-Transmittal"

    date_str = extract_cover_date_bulletproof(full_text)

    subject = ""
    m_sub = re.search(r'(?:SUBJECT|Subject)\s*[:\s]+([^\n\r]+)', full_text, re.IGNORECASE)
    if m_sub:
        raw_sub = m_sub.group(1).strip()
        raw_sub = re.split(r'(?:Reference|Dear|Attn|Attachment|Document/Drawing|Sheet:)', raw_sub, flags=re.IGNORECASE)[0].strip()
        subject = raw_sub
    if not subject:
        subject = "Official Transmittal"

    lines = full_text.splitlines()
    purpose_markers = {'IFA', 'IFI', 'APD', 'AB', 'RI', 'END', '- END -', '-END-', '-'}
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        m_doc = re.search(r'(QT1-[0-9L]-[A-Za-z0-9_-]+)', line)
        if m_doc and "EVNPMB2" not in m_doc.group(1) and "QT1-HDH" not in m_doc.group(1):
            raw_code = m_doc.group(1).strip()
            norm_code = normalize_doc_code(raw_code)
            after = line[m_doc.end():].strip()
            
            words = after.split()
            while words and words[-1].upper() in purpose_markers:
                words.pop()
            while words and words[-1] in ['-', 'END']:
                words.pop()
                
            rev = "1"
            if len(words) >= 2 and words[0].isdigit() and (len(words[1]) <= 3 and words[1].isalnum() and not words[1].isdigit()):
                words.pop(0)
                rev = words.pop(0)
            elif len(words) >= 1 and len(words[0]) <= 3 and words[0].isalnum():
                rev = words.pop(0)
            elif len(words) >= 1 and re.match(r'^Rev\.?([A-Za-z0-9]+)$', words[0], re.IGNORECASE):
                m_r = re.match(r'^Rev\.?([A-Za-z0-9]+)$', words[0], re.IGNORECASE)
                rev = m_r.group(1)
                words.pop(0)
                
            title_parts = [ " ".join(words).strip() ]
            
            j = i + 1
            while j < len(lines):
                next_line = lines[j].strip()
                if re.search(r'(QT1-[0-9L]-[A-Za-z0-9_-]+)', next_line) or re.match(r'^\d{1,2}\s+', next_line) or "Purpose:" in next_line:
                    break
                if next_line and not any(k in next_line for k in ["Sheet:", "Date:", "Page", "END -"]):
                    clean_next = " ".join([w for w in next_line.split() if w.upper() not in purpose_markers])
                    if clean_next:
                        title_parts.append(clean_next)
                j += 1
                
            combined_title = clean_document_title(" ".join(title_parts))
            
            if norm_code and combined_title:
                docs[norm_code] = {
                    'doc_no': norm_code,
                    'rev': rev,
                    'title': combined_title
                }
        i += 1

    return {
        'tr_no': tr_no,
        'date_str': date_str if date_str else datetime.now().strftime("%d/%m/%Y"),
        'subject': subject,
        'docs': docs
    }

def find_template(prefix):
    search_dirs = [BASE_DIR, os.getcwd(), r"G:\Shared drives\Test auto"]
    for d in search_dirs:
        if not os.path.exists(d): continue
        for f in os.listdir(d):
            if f.lower().endswith(".docx"):
                clean_f = f.lower().replace("đ", "d").replace("ð", "d")
                clean_p = prefix.lower().replace("đ", "d").replace("ð", "d")
                if clean_p in clean_f and not f.startswith("~$"):
                    return os.path.join(d, f)
    return None

def find_car_template():
    search_dirs = [BASE_DIR, os.getcwd(), r"G:\Shared drives\Test auto"]
    for d in search_dirs:
        if not os.path.exists(d): continue
        for f in os.listdir(d):
            if f.lower().endswith(".xlsx") and not f.startswith("~$"):
                if any(k in f.lower() for k in ["template", "car", "00xx", "xxxx"]):
                    return os.path.join(d, f)
    return None

def create_fresh_car_wb():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ROUND 1 CaRs"
    ws.views.sheetView[0].showGridLines = True
    col_widths = {'A': 4.5, 'B': 1.5, 'C': 16.0, 'D': 1.5, 'E': 20.0, 'F': 1.5, 'G': 48.0, 'H': 1.5, 'I': 58.0, 'J': 1.5, 'K': 4.5, 'L': 14.0}
    for col, width in col_widths.items(): ws.column_dimensions[col].width = width
    dot = Side(border_style="dotted", color="808080")
    dotted_box = Border(top=dot, left=dot, right=dot, bottom=dot)
    f_title = Font(name="Times New Roman", size=16, bold=True)
    f_car_no = Font(name="Times New Roman", size=16, bold=True, color="FFFF0000")
    f_label, f_header = Font(name="Arial", size=10, bold=False), Font(name="Arial", size=10, bold=True)
    f_status = Font(name="Arial", size=10, bold=True, color="FFFF0000")
    al_center, al_left = Alignment(horizontal="center", vertical="center"), Alignment(horizontal="left", vertical="center")
    
    ws.merge_cells("A1:D1")
    ws['A1'], ws['A1'].font, ws['A1'].alignment = "CaR Nr.", f_title, al_left
    ws['E1'].font, ws['E1'].alignment = f_car_no, al_left

    labels = [(2, "i", "Sender of Document"), (3, "ii", "Covering Letter No & Date"), (4, "iii", "Doc Ident No & Rev"), (5, "iv", "Title of Document"), (6, "v", "Date Doc Rec'd by EVN/Consultant"), (7, "vi", "Reviewed by & date completed"), (8, "vii", "Date & No CaRs is issued to PMB2"), (9, "viii", "Date & No CaRs is issued to Contractor")]
    for r, num_rom, text in labels:
        ws.cell(row=r, column=1, value=num_rom).font, ws.cell(row=r, column=1).alignment = f_label, al_center
        ws.merge_cells(f"C{r}:E{r}")
        ws.cell(row=r, column=3, value=text).font, ws.cell(row=r, column=3).alignment = f_label, al_left
        ws.cell(row=r, column=7).font, ws.cell(row=r, column=7).alignment = f_header, al_left

    ws.cell(row=10, column=3, value="Contract /").font = f_header
    ws.cell(row=10, column=5, value="Doc Cls/Page").font = f_header
    ws.cell(row=10, column=9, value="CONTRACTOR's").font = f_header
    ws.cell(row=10, column=9).alignment = al_left

    headers = [(1, "No"), (3, "Spec Ref"), (5, "Drg Ref etc"), (7, "ROUND 1 COMMENTS"), (9, "ROUND 1 REPLIES")]
    for c, h_text in headers:
        ws.cell(row=11, column=c, value=h_text).font = f_header
        ws.cell(row=11, column=c).alignment = al_center if c == 1 else al_left

    ws.cell(row=13, column=7, value="STATUS AFTER ROUND 1 COMMENTS:").font = f_status
    ws.cell(row=14, column=7, value="APPROVED (APP)").font = f_status
    ws.cell(row=15, column=7, value="General Comments").font = f_header
    ws.cell(row=16, column=1, value=1).font, ws.cell(row=16, column=1).alignment = f_label, al_center
    ws.cell(row=16, column=7, value="No comment.").font = f_label
    
    for r in range(17, 27):
        ws.cell(row=r, column=1, value=r - 15).font, ws.cell(row=r, column=1).alignment = f_label, al_center
    ws.cell(row=27, column=7, value="Detailed Comments").font = f_header
    for r in range(28, 36):
        ws.cell(row=r, column=1, value=r - 27).font, ws.cell(row=r, column=1).alignment = f_label, al_center
    for r in range(1, 36):
        for c in range(1, 13):
            if r >= 11: ws.cell(row=r, column=c).border = dotted_box
    return wb

def replace_value_after_keyword(paragraph, keyword, new_text):
    if keyword in paragraph.text:
        if not paragraph.runs:
            paragraph.text = f"{keyword} {new_text}"
            return True
        kw_font = paragraph.runs[0].font
        kw_name, kw_size, kw_bold, kw_italic = kw_font.name, kw_font.size, kw_font.bold, kw_italic
        val_font = paragraph.runs[-1].font
        val_name, val_size, val_bold, val_italic = val_font.name, val_font.size, val_bold, val_italic
        
        paragraph.clear()
        r1 = paragraph.add_run(keyword)
        r1.font.name, r1.font.size, r1.font.bold, r1.font.italic = kw_name, kw_size, kw_bold, kw_italic
        r2 = paragraph.add_run(" " + new_text)
        r2.font.name, r2.font.size, r2.font.bold, r2.font.italic = val_name, val_size, val_bold, val_italic
        return True
    return False

def update_date_paragraph(paragraph, new_date_str):
    val_font = paragraph.runs[-1].font
    val_name, val_size, val_bold, val_italic = val_font.name, val_font.size, val_bold, val_italic
    paragraph.clear()
    paragraph.add_run("                                ")
    r2 = paragraph.add_run(new_date_str)
    r2.font.name, r2.font.size, r2.font.bold, r2.font.italic = val_name, val_size, val_bold, val_italic

def sync_footer_content(doc, qti_no, subject):
    new_text = f"{qti_no} {subject}"
    for section in doc.sections:
        if section.footer:
            for p in section.footer.paragraphs:
                if p.text.strip():
                    if len(p.runs) > 0:
                        f_font = p.runs[0].font
                        p.clear()
                        r = p.add_run(new_text)
                        r.font.name, r.font.size = f_font.name, f_font.size
                        r.font.color.rgb = f_font.color.rgb if f_font.color else RGBColor(128, 128, 128)
                    else:
                        p.text = new_text

def remove_temp_file(path: str):
    try:
        if os.path.exists(path): os.remove(path)
    except: pass

@app.post("/api/export-zip-full")
async def export_zip_full(
    files: list[UploadFile] = File(default=[]),
    cover_file: UploadFile = File(None),
    metadata: str = Form("{}"),
    mode: str = Form("auto_bump"),
    qti_no: str = Form("QTI-PECC1-PMB2-26-10-155"),
    start_car_no: str = Form("9905"),
    custom_date: str = Form(""),
    default_status: str = Form("APP")
):
    temp_zip = tempfile.NamedTemporaryFile(delete=False, suffix=".zip")
    temp_zip_path = temp_zip.name
    temp_zip.close()

    try:
        car_metadata = json.loads(metadata) if metadata else {}
        now = datetime.now()
        current_date_str = now.strftime("%d %b %Y")
        current_date_en = now.strftime("%B %d, %Y")
        
        cover_data = {
            'tr_no': 'QT1-Document',
            'date_str': now.strftime("%d/%m/%Y"),
            'subject': 'Official Transmittal',
            'docs': {}
        }
        
        if cover_file is not None:
            cover_bytes = await cover_file.read()
            if cover_file.filename.lower().endswith('.pdf'):
                cover_data = extract_cover_metadata(cover_bytes, cover_file.filename)
            del cover_bytes

        if custom_date and custom_date.strip():
            cover_data['date_str'] = custom_date.strip()

        vi_template_path = find_template("TV")
        en_template_path = find_template("QTI-PECC1")
        if not vi_template_path or not en_template_path: 
            raise FileNotFoundError("Thiếu file Word mẫu (TVÐ1-NÐNLM.docx / QTI-PECC1.docx) trong máy chủ.")

        processed_car_list = []
        start_num = int(re.sub(r'\D', '', str(start_car_no)) or "9905")

        with zipfile.ZipFile(temp_zip_path, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            
            # 1. TỰ ĐỘNG TẠO CAR
            if mode == "auto_generate_cars":
                docs_to_generate = list(cover_data['docs'].values()) if cover_data['docs'] else [{'doc_no': cover_data['tr_no'], 'rev': 'A', 'title': cover_data['subject']}]
                status_full = "APPROVED (APP)" if default_status == "APP" else "RETURN FOR CORRECTIONS (RFC)"
                if default_status == "AWC": status_full = "APPROVED WITH COMMENTS (AWC)"
                elif default_status == "RJ": status_full = "REJECTED (RJ)"

                car_template_path = find_car_template()

                for doc_idx, doc_item in enumerate(docs_to_generate):
                    cur_num = start_num + doc_idx
                    car_num_str = f"{cur_num:04d}" if cur_num < 10000 else f"{cur_num}"
                    car_base = f"{car_num_str} CaRs"

                    wb = openpyxl.load_workbook(car_template_path) if car_template_path else create_fresh_car_wb()
                    for ws_item in wb.worksheets:
                        if hasattr(ws_item, 'sheet_properties') and ws_item.sheet_properties and hasattr(ws_item.sheet_properties, 'pageSetUpPr'):
                            ws_item.sheet_properties.pageSetUpPr.fitToPage = None

                    ws = wb['ROUND 1 CaRs'] if 'ROUND 1 CaRs' in wb.sheetnames else wb.active
                    ws['E1'].value = f"{car_num_str}-R1"
                    ws['E1'].font = Font(name="Times New Roman", size=16, bold=True, color="FFFF0000")
                    ws['E1'].fill = PatternFill(fill_type=None)

                    ws['G2'].value = "HYUNDAI - MITSUBISHI - CC1"
                    ws['G3'].value = cover_data['tr_no']
                    ws['G4'].value = f"{doc_item['doc_no']} rev.{doc_item['rev']}"
                    ws['G5'].value = doc_item['title']
                    ws['G6'].value = cover_data['date_str']

                    ws['G13'].value = "STATUS AFTER ROUND 1 COMMENTS:"
                    ws['G14'].value = status_full
                    if default_status == "APP": ws['G16'].value = "No comment."

                    new_excel_name = f"{car_base} R1 {current_date_str} {doc_item['doc_no']} {clean_safe_filename(doc_item['title'])} rev.{doc_item['rev']}.xlsx"
                    excel_out = io.BytesIO()
                    wb.save(excel_out)
                    zip_file.writestr(new_excel_name, excel_out.getvalue())
                    del wb, excel_out

                    processed_car_list.append({
                        'doc_no': doc_item['doc_no'], 'rev': doc_item['rev'], 'title': doc_item['title'],
                        'status': status_full, 'car_label': f"{car_base} R1", 'assigned_car_name': new_excel_name
                    })

            # 2. XỬ LÝ HÀNG LOẠT FILE TẢI LÊN
            elif mode != "no_car_letter" and files:
                for file_idx, file in enumerate(files):
                    try:
                        content = await file.read()
                        wb = openpyxl.load_workbook(io.BytesIO(content))
                        
                        for ws_item in wb.worksheets:
                            if hasattr(ws_item, 'sheet_properties') and ws_item.sheet_properties and hasattr(ws_item.sheet_properties, 'pageSetUpPr'):
                                ws_item.sheet_properties.pageSetUpPr.fitToPage = None

                        max_r = 1
                        target_idx = 0
                        for idx, s_name in enumerate(wb.sheetnames):
                            m = re.search(r'round\s*(\d+)', s_name, re.IGNORECASE)
                            if m and int(m.group(1)) >= max_r:
                                max_r = int(m.group(1))
                                target_idx = idx

                        source_ws = wb.worksheets[target_idx]
                        
                        prev_com_col, prev_rep_col = None, None
                        for c in range(1, source_ws.max_column + 1):
                            val_h = str(source_ws.cell(row=11, column=c).value or "").upper()
                            if f"ROUND {max_r} COMMENTS" in val_h: prev_com_col = c
                            elif f"ROUND {max_r} REPLIES" in val_h: prev_rep_col = c
                        if not prev_com_col: prev_com_col = 7

                        # BÓC TÁCH MÃ DOC VÀ TÊN (3-TIER HIERARCHY CHỐNG SAI LỆCH)
                        detected_doc_code = ""
                        for c in range(3, source_ws.max_column + 1):
                            v_c = str(source_ws.cell(row=4, column=c).value or "")
                            if "QT1-" in v_c:
                                detected_doc_code = normalize_doc_code(v_c)
                                break
                        if not detected_doc_code: 
                            detected_doc_code = normalize_doc_code(file.filename)

                        cover_info = cover_data['docs'].get(detected_doc_code, {})
                        final_doc_no = cover_info.get('doc_no', detected_doc_code if detected_doc_code else "QT1-Document")
                        final_rev = cover_info.get('rev', '1')

                        # Ưu tiên 1: Lấy từ Transmittal nếu có
                        if cover_info.get('title') and cover_info['title'].lower() != cover_data['subject'].lower():
                            final_title = cover_info['title']
                        else:
                            # Ưu tiên 2: Lấy từ Tên file CaR kỹ sư tải lên
                            title_from_fn = extract_title_from_filename(file.filename)
                            if title_from_fn:
                                final_title = title_from_fn
                            else:
                                # Ưu tiên 3: Lấy từ ô G5 của Round cũ đã làm sạch
                                r5_val = clean_document_title(str(source_ws.cell(row=5, column=prev_com_col).value or ""))
                                final_title = r5_val if r5_val else cover_data['subject']

                        cur_num = start_num + file_idx
                        car_num_str = f"{cur_num:04d}" if cur_num < 10000 else f"{cur_num}"
                        car_base = f"{car_num_str} CaRs"

                        # =======================================================
                        # CHẾ ĐỘ: CẤP SỐ CHO CAR KỸ SƯ
                        # =======================================================
                        if mode == "number_filled_cars":
                            detected_status_raw = "APPROVED (APP)"
                            found_rfc, found_awc, found_rj = False, False, False
                            for r in range(11, 16):
                                v_st = str(source_ws.cell(row=r, column=prev_com_col).value or "").strip().upper()
                                if "RETURN FOR CORRECTIONS" in v_st or "(RFC)" in v_st or v_st == "RFC": found_rfc = True
                                elif "APPROVED WITH COMMENTS" in v_st or "(AWC)" in v_st or v_st == "AWC": found_awc = True
                                elif "REJECTED" in v_st or "(RJ)" in v_st or v_st == "RJ": found_rj = True
                            if found_rfc: detected_status_raw = "RETURN FOR CORRECTIONS (RFC)"
                            elif found_awc: detected_status_raw = "APPROVED WITH COMMENTS (AWC)"
                            elif found_rj: detected_status_raw = "REJECTED (RJ)"
                            
                            found_cell_r1 = False
                            for c in range(2, 10):
                                val_r1 = str(source_ws.cell(row=1, column=c).value or "").strip()
                                if re.search(r'(00xx|xxxx|Xxxx)', val_r1, re.IGNORECASE):
                                    source_ws.cell(row=1, column=c).value = f"{car_base} - R{max_r}"
                                    source_ws.cell(row=1, column=c).font = Font(name="Times New Roman", size=16, bold=True, color="FFFF0000")
                                    source_ws.cell(row=1, column=c).fill = PatternFill(fill_type=None)
                                    found_cell_r1 = True
                                    break
                            if not found_cell_r1:
                                source_ws.cell(row=1, column=5).value = f"{car_base} - R{max_r}"
                                source_ws.cell(row=1, column=5).font = Font(name="Times New Roman", size=16, bold=True, color="FFFF0000")
                                source_ws.cell(row=1, column=5).fill = PatternFill(fill_type=None)

                            source_ws.cell(row=4, column=prev_com_col).value = f"{final_doc_no} rev.{final_rev}"
                            source_ws.cell(row=5, column=prev_com_col).value = final_title

                            new_excel_name = f"{car_base} R{max_r} {current_date_str} {final_doc_no} {clean_safe_filename(final_title)} rev.{final_rev}.xlsx"
                            excel_out = io.BytesIO()
                            wb.save(excel_out)
                            zip_file.writestr(new_excel_name, excel_out.getvalue())
                            del wb, excel_out

                            processed_car_list.append({
                                'doc_no': final_doc_no, 'rev': final_rev, 'title': final_title,
                                'status': detected_status_raw, 'car_label': f"{car_base} R{max_r}", 'assigned_car_name': new_excel_name
                            })

                        # =======================================================
                        # CHẾ ĐỘ: NÂNG ROUND (R+1) SIÊU TỐC
                        # =======================================================
                        elif mode == "auto_bump":
                            selected_status = car_metadata.get(file.filename, "APP")
                            new_r_num = max_r + 1
                            new_ws = wb.copy_worksheet(source_ws)
                            new_ws.title = f"ROUND {new_r_num} CaRs"
                            wb._sheets.remove(new_ws)
                            wb._sheets.insert(target_idx + 1, new_ws)

                            new_com_col = (prev_rep_col + 1) if prev_rep_col else (new_ws.max_column + 1)
                            new_rep_col = new_com_col + 2

                            for c in range(2, 10):
                                v = str(new_ws.cell(row=1, column=c).value or "")
                                if re.search(r'-\s*R\d+', v, re.IGNORECASE) or (c == 5 and v):
                                    new_ws.cell(row=1, column=c).value = re.sub(r'R\d+', f'R{new_r_num}', v, flags=re.IGNORECASE)
                                    break

                            new_ws.cell(row=2, column=new_com_col).value = "EVNPMB2"
                            new_ws.cell(row=3, column=new_com_col).value = cover_data['tr_no']
                            new_ws.cell(row=4, column=new_com_col).value = f"{final_doc_no} rev.{final_rev}"
                            new_ws.cell(row=5, column=new_com_col).value = final_title
                            new_ws.cell(row=10, column=new_rep_col).value = "CONTRACTOR's"
                            new_ws.cell(row=11, column=new_com_col).value = f"ROUND {new_r_num} COMMENTS"
                            new_ws.cell(row=11, column=new_rep_col).value = f"ROUND {new_r_num} REPLIES"

                            status_full = "APPROVED (APP)" if selected_status == "APP" else "RETURN FOR CORRECTIONS (RFC)"
                            if selected_status == "AWC": status_full = "APPROVED WITH COMMENTS (AWC)"
                            elif selected_status == "RJ": status_full = "REJECTED (RJ)"

                            new_ws.cell(row=12, column=new_com_col).value = f"STATUS AFTER ROUND {new_r_num} COMMENTS:"
                            new_ws.cell(row=13, column=new_com_col).value = status_full
                            red_bold = Font(name="Arial", size=10, bold=True, color="FFFF0000")
                            new_ws.cell(row=12, column=new_com_col).font, new_ws.cell(row=13, column=new_com_col).font = red_bold, red_bold
                            new_ws.column_dimensions[openpyxl.utils.get_column_letter(new_com_col)].width = 38
                            new_ws.column_dimensions[openpyxl.utils.get_column_letter(new_rep_col)].width = 38

                            # Duyệt tối ưu các hàng nhận xét
                            limit_row = min(new_ws.max_row, 100)
                            for r in range(14, limit_row + 1):
                                v_prev_c = str(new_ws.cell(row=r, column=prev_com_col).value or "").strip()
                                if v_prev_c and not ("general comment" in v_prev_c.lower() or "detailed comment" in v_prev_c.lower()):
                                    if selected_status == "APP": 
                                        new_ws.cell(row=r, column=new_com_col).value = "Noted."

                            new_excel_name = f"{car_base} R{new_r_num} {current_date_str} {final_doc_no} {clean_safe_filename(final_title)} rev.{final_rev}.xlsx"
                            excel_out = io.BytesIO()
                            wb.save(excel_out)
                            zip_file.writestr(new_excel_name, excel_out.getvalue())
                            del wb, excel_out

                            processed_car_list.append({
                                'doc_no': final_doc_no, 'rev': final_rev, 'title': final_title,
                                'status': status_full, 'car_label': f"{car_base} R{new_r_num}", 'assigned_car_name': new_excel_name
                            })

                        # =======================================================
                        # CHẾ ĐỘ: TRÍCH XUẤT CAR CÓ SẴN
                        # =======================================================
                        elif mode == "pre_prepared":
                            status_full = "APPROVED (APP)"
                            for r in [13, 14, 12]:
                                st_val = str(source_ws.cell(row=r, column=prev_com_col).value or "").upper()
                                if "RFC" in st_val: status_full = "RETURN FOR CORRECTIONS (RFC)"
                                elif "AWC" in st_val: status_full = "APPROVED WITH COMMENTS (AWC)"
                                elif "RJ" in st_val: status_full = "REJECTED (RJ)"

                            car_base = re.search(r'(\d{4,}\s*CaRs?)', file.filename, re.IGNORECASE)
                            car_lbl = car_base.group(1) if car_base else "CaRs"
                            
                            zip_file.writestr(clean_safe_filename(file.filename), content)
                            processed_car_list.append({
                                'doc_no': final_doc_no, 'rev': final_rev, 'title': final_title,
                                'status': status_full, 'car_label': f"{car_lbl} R{max_r}", 'assigned_car_name': file.filename
                            })
                        
                        del content
                        gc.collect()

                    except Exception as fe:
                        print(f"Lỗi tệp {file.filename}: {fe}")

            # 3. TRẢ LỜI THƯ THUẦN TÚY (KHÔNG KÈM CAR)
            elif mode == "no_car_letter":
                if cover_data['docs']:
                    for doc_k, d_info in reversed(list(cover_data['docs'].items())):
                        processed_car_list.append({
                            'doc_no': d_info['doc_no'], 'rev': d_info['rev'], 'title': d_info['title'],
                            'status': 'REVIEWED (NO CAR)', 'car_label': 'N/A', 'assigned_car_name': 'N/A'
                        })
                else:
                    processed_car_list.append({
                        'doc_no': cover_data['tr_no'], 'rev': '0', 'title': cover_data['subject'],
                        'status': 'REVIEWED (NO CAR)', 'car_label': 'N/A', 'assigned_car_name': 'N/A'
                    })

            # =========================================================================
            # ĐỒNG BỘ 2 THƯ WORD COVER
            # =========================================================================
            safe_subj_name = clean_safe_filename(cover_data['subject'])

            doc_vi = Document(vi_template_path)
            for tbl in doc_vi.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            if "ngày" in p.text and "tháng" in p.text:
                                p.runs[0].text = f"Hà Nội, ngày {now.day} tháng {now.month} năm {now.year}"
                                for r in p.runs[1:]: r.text = ""
                            if "V/v" in p.text: replace_value_after_keyword(p, "V/v", cover_data['subject'])

            for p in doc_vi.paragraphs:
                if "Căn cứ" in p.text or "Công ty Cổ phần" in p.text: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if "Căn cứ công văn số" in p.text:
                    tgt = f"{cover_data['tr_no']} ngày {cover_data['date_str']} của Nhà thầu EPC gói 15 về {cover_data['subject']}."
                    if not replace_value_after_keyword(p, "Căn cứ công văn số", tgt): p.text = f"Căn cứ công văn số {tgt}"

            vi_io = io.BytesIO()
            doc_vi.save(vi_io)
            zip_file.writestr(f"TVÐ1-NÐNLM {safe_subj_name}.docx", vi_io.getvalue())

            doc_en = Document(en_template_path)
            doc_type_en = "Letter" if "-L-" in cover_data['tr_no'].upper() else "Transmittal"

            for tbl in doc_en.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            if "Ref. No.:" in p.text: replace_value_after_keyword(p, "Ref. No.:", cover_data['tr_no'])
                            elif "October" in p.text or "202" in p.text: update_date_paragraph(p, current_date_en)
                            elif "Identification Number." in p.text: replace_value_after_keyword(p, "Identification Number.", qti_no.strip())
                            elif "Subject:" in p.text: replace_value_after_keyword(p, "Subject:", cover_data['subject'])

            sync_footer_content(doc_en, qti_no.strip(), cover_data['subject'])

            ref_idx, yours_idx = None, None
            for idx, p in enumerate(doc_en.paragraphs):
                if "This has reference to the" in p.text:
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    f_name = p.runs[0].font.name if (p.runs and p.runs[0].font.name) else "Times New Roman"
                    f_size = p.runs[0].font.size if (p.runs and p.runs[0].font.size) else Pt(13)
                    p.clear()
                    r = p.add_run(f"This has reference to the {doc_type_en} {cover_data['tr_no']}, we would like to comment as follows:")
                    r.font.name, r.font.size = f_name, f_size
                    ref_idx = idx
                elif "Yours sincerely" in p.text:
                    p.alignment, yours_idx = WD_ALIGN_PARAGRAPH.LEFT, idx
                    break

            if ref_idx is not None and yours_idx is not None:
                for p in [doc_en.paragraphs[i] for i in range(ref_idx + 1, yours_idx)]:
                    p._element.getparent().remove(p._element)

            target_p = None
            for p in doc_en.paragraphs:
                if "Yours sincerely" in p.text:
                    target_p = p
                    break

            if target_p is not None:
                if mode == "no_car_letter":
                    items_iter = reversed(list(cover_data['docs'].values())) if cover_data['docs'] else [{"doc_no": cover_data['tr_no'], "title": cover_data['subject']}]
                    for it in items_iter:
                        new_p = target_p.insert_paragraph_before("")
                        new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        new_p.paragraph_format.left_indent, new_p.paragraph_format.first_line_indent = Inches(0.35), Inches(-0.35)
                        new_p.paragraph_format.space_before, new_p.paragraph_format.space_after = 0, Pt(4)
                        if cover_data['docs']: txt = f"-\tThe document {it['doc_no']} {it['title']} Rev.{it['rev']} has been reviewed and found acceptable with no further comments."
                        else: txt = f"-\tWe have reviewed the submitted transmittal/documents and have no objection."
                        run = new_p.add_run(txt)
                        run.font.name, run.font.size = 'Times New Roman', Pt(13)
                else:
                    for item in reversed(processed_car_list):
                        new_p = target_p.insert_paragraph_before("")
                        new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        new_p.paragraph_format.left_indent, new_p.paragraph_format.first_line_indent = Inches(0.35), Inches(-0.35)
                        new_p.paragraph_format.space_before, new_p.paragraph_format.space_after = 0, Pt(4)
                        run = new_p.add_run(f"-\tThe document {item['doc_no']} {item['title']} Rev.{item['rev']} is {item['status']} indicated as per {item['car_label']}.")
                        run.font.name, run.font.size = 'Times New Roman', Pt(13)

            en_io = io.BytesIO()
            doc_en.save(en_io)
            zip_file.writestr(f"{clean_safe_filename(qti_no)} {safe_subj_name}.docx", en_io.getvalue())
            
        headers = {"Content-Disposition": f'attachment; filename="{clean_safe_filename(qti_no)}.zip"'}
        try:
            items_json = json.dumps([{"doc_no": i['doc_no'], "assigned_car": i['assigned_car_name'], "status": i['status']} for i in processed_car_list], ensure_ascii=False)
            b64_str = base64.b64encode(items_json.encode('utf-8')).decode('ascii')
            if len(b64_str) < 6000: headers["X-Processed-Items"] = b64_str
        except: pass

        return FileResponse(path=temp_zip_path, media_type="application/zip", filename=f"{clean_safe_filename(qti_no)}.zip", headers=headers, background=BackgroundTask(remove_temp_file, temp_zip_path))

    except Exception as ge:
        remove_temp_file(temp_zip_path)
        print("Lỗi hệ thống:", ge)
        traceback.print_exc()
        return Response(content=f"Error: {str(ge)}", status_code=400)

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
